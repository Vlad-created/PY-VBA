import win32com.client as win32
import os
import sys
import shutil
import time
import subprocess
import pandas as pd
from openpyxl import load_workbook, Workbook

# ============================================================
# CONFIGURATION
# ============================================================

SENDER_NAME = "Rumen I. Dimitrov"
SUBJECT_MATCH = "Дигитално разписани бланки"

destination_path = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\Дигитално разписани - получени файлове\TTl registrations - to upload.xlsx"
refresh_file_path = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\TTl registrations (002).xlsx"

access_db_path = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\Mandarin.accdb"
access_macro = "Macro1"

input_file = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\dealers\Prima_Sales_Registration.xlsx"
output_folder = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\dealers"

groups = [
    "KANDYLTD", "VARDAROV", "YONI-90", "MOBIMAKSLTD",
    "POSTBRANCH", "HAIKOVLTD", "ET_BOMA", "PREPAID TEAM", "EasyPay", "Fast Pay"
]

sheet_name = "Prima_Sales_Registration_export"

# ============================================================
# STEP 1 — DOWNLOAD EMAIL ATTACHMENT
# ============================================================

print("STEP 1: Checking Outlook for email...")

outlook = win32.Dispatch("Outlook.Application").GetNamespace("MAPI")

try:
    inbox = outlook.GetDefaultFolder(6)
except Exception as e:
    print("ERROR: Cannot access Outlook Inbox:", e)
    sys.exit(1)

messages = inbox.Items
messages = messages.Restrict("[MessageClass]='IPM.Note'")
messages.Sort("ReceivedTime", True)

target_email = None
for msg in messages:
    try:
        if SENDER_NAME.lower() in msg.SenderName.lower() and SUBJECT_MATCH.lower() in msg.Subject.lower():
            target_email = msg
            break
    except:
        continue

if not target_email:
    print("No matching email found.")
    sys.exit(0)

print(f"Found email from: {target_email.SenderName}")
print(f"Subject: {target_email.Subject}")

excel_attachment = None
for att in target_email.Attachments:
    if att.FileName.lower().endswith((".xlsx", ".xlsm", ".xls")):
        excel_attachment = att
        break

if not excel_attachment:
    print("No Excel attachment found.")
    sys.exit(0)

temp_path = os.path.join(os.environ["TEMP"], "temp_attachment.xlsx")
excel_attachment.SaveAsFile(temp_path)

if not os.path.exists(temp_path):
    print("ERROR: Attachment failed to save.")
    sys.exit(1)

print("Attachment downloaded successfully.")

# Check if destination file is locked
try:
    with open(destination_path, "a"):
        pass
except PermissionError:
    print("Destination file is OPEN or LOCKED. Close it and try again.")
    sys.exit(1)

try:
    shutil.copyfile(temp_path, destination_path)
    os.remove(temp_path)
except Exception as e:
    print("ERROR copying file:", e)
    sys.exit(1)

print(f"SUCCESS: Attachment saved to:\n{destination_path}")

# ============================================================
# STEP 2 — REFRESH EXCEL QUERIES (RELIABLE VERSION)
# ============================================================

print("\nSTEP 2: Refreshing Excel queries...")

excel = win32.Dispatch("Excel.Application")
excel.Visible = False  # keep hidden

try:
    wb = excel.Workbooks.Open(refresh_file_path)
except Exception as e:
    print("ERROR opening Excel file:", e)
    excel.Quit()
    sys.exit(1)

# Disable background refresh
for conn in wb.Connections:
    try:
        conn.OLEDBConnection.BackgroundQuery = False
    except:
        pass

print("Starting refresh...")
wb.RefreshAll()

# --- Function to check if Excel is still refreshing ---
def is_refreshing_all(wb):
    # Check connections
    for conn in wb.Connections:
        try:
            if conn.Refreshing:
                return True
        except:
            pass

    # Check if Excel is busy with background tasks
    try:
        if excel.CalculateState != 0:  # 0 = xlDone
            return True
    except:
        pass

    return False

# --- Wait loop ---
print("Waiting for Excel to finish refreshing...")
while is_refreshing_all(wb):
    time.sleep(1)

# --- Extra safety wait ---
print("Finalizing refresh (extra wait 45 seconds)...")
time.sleep(45)

# Save and close
wb.Save()
wb.Close()
excel.Quit()

print("Excel refreshed, saved, and closed successfully.")

# ============================================================
# STEP 3 — RUN ACCESS MACRO
# ============================================================

print("\nSTEP 3: Running Access macro...")

subprocess.call("taskkill /f /im msaccess.exe", shell=True)

access = win32.Dispatch("Access.Application")
access.Visible = False

access.OpenCurrentDatabase(access_db_path)
access.DoCmd.RunMacro(access_macro)
access.CloseCurrentDatabase()
access.Quit()

print("Access macro completed and Access closed.")

# ============================================================
# STEP 4 — FORMAT & SPLIT EXCEL FILES
# ============================================================

print("\nSTEP 4: Formatting and splitting Excel files...")

wb = load_workbook(input_file)
ws = wb[sheet_name]

date_format = "yyyy-mm-dd"
text_format = "@"

for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=1).number_format = date_format
    ws.cell(row=row, column=8).number_format = date_format

    cell_b = ws.cell(row=row, column=2)
    cell_c = ws.cell(row=row, column=3)

    if cell_b.value is not None:
        cell_b.value = str(cell_b.value)
    if cell_c.value is not None:
        cell_c.value = str(cell_c.value)

    cell_b.number_format = text_format
    cell_c.number_format = text_format

wb.save(input_file)
print("Main file formatted.")

df = pd.read_excel(input_file, sheet_name=sheet_name)
column_name = df.columns[5]

for group in groups:
    print(f"Processing group: {group}")

    filtered_df = df[df[column_name].astype(str).str.contains(group, case=False, na=False)]

    if filtered_df.empty:
        print(f"  No rows for {group}, skipping.")
        continue

    output_file = os.path.join(output_folder, f"Prima Sales Registration_{group}.xlsx")

    new_wb = Workbook()
    ws_new = new_wb.active
    ws_new.title = "Data"

    ws_new.append(list(filtered_df.columns))

    for _, row in filtered_df.iterrows():
        ws_new.append(list(row))

    for r in range(2, ws_new.max_row + 1):
        ws_new.cell(row=r, column=1).number_format = date_format
        ws_new.cell(row=r, column=8).number_format = date_format

        cell_b = ws_new.cell(row=r, column=2)
        cell_c = ws_new.cell(row=r, column=3)

        if cell_b.value is not None:
            cell_b.value = str(cell_b.value)
        if cell_c.value is not None:
            cell_c.value = str(cell_c.value)

        cell_b.number_format = text_format
        cell_c.number_format = text_format

    new_wb.save(output_file)
    print(f"  Saved: {output_file}")

print("\nALL STEPS COMPLETED SUCCESSFULLY.")
