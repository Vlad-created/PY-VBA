import openpyxl
import xlwings as xw
import os
import datetime

folder = r"\\virgo\Work\OpControl\2016\Bonuses\1_MassCurrMonth\Targets"

current_month = datetime.datetime.now().strftime("%B")

# -----------------------------
# Locate SOURCE file
# -----------------------------
src_path = None
for filename in os.listdir(folder):
    if (
        filename.startswith("Individual ")
        and filename.endswith("Final_Merged.xlsx")
        and current_month in filename
    ):
        src_path = os.path.join(folder, filename)
        break

if not src_path:
    raise FileNotFoundError("Source file not found.")

# -----------------------------
# Locate DESTINATION file
# -----------------------------
dst_path = None
for filename in os.listdir(folder):
    if (
        filename.startswith("Individual_targets_")
        and filename.endswith("_Total - Copy.xlsx")
        and current_month in filename
    ):
        dst_path = os.path.join(folder, filename)
        break

if not dst_path:
    raise FileNotFoundError("Destination file not found.")

print(f"Source: {src_path}")
print(f"Destination: {dst_path}")

# -----------------------------
# Load SOURCE workbook
# -----------------------------
src_wb = openpyxl.load_workbook(src_path, data_only=True)

if "APPENDED" not in src_wb.sheetnames:
    raise ValueError("Sheet 'APPENDED' not found.")

src_ws = src_wb["APPENDED"]

last_row = src_ws.max_row

# Read columns B–I and K
data = []
col_map = ["B", "C", "D", "E", "F", "G", "H", "I", "K"]

for row in range(2, last_row + 1):
    row_values = [src_ws[f"{col}{row}"].value for col in col_map]
    safe_row = [str(v) if v is not None else "" for v in row_values]
    data.append(safe_row)

print(f"Loaded {len(data)} rows from source.")

# -----------------------------
# Open DESTINATION workbook
# -----------------------------
app = xw.App(visible=True, add_book=False)
dst_wb = app.books.open(dst_path, update_links=True)
dst_ws = dst_wb.sheets["individual"]

# -----------------------------
# FORMAT COLUMNS
# -----------------------------

# TEXT columns (preserve leading zeros)
text_columns = ["D", "E", "F", "H", "I", "J", "K"]
for col in text_columns:
    dst_ws.range(f"{col}:{col}").api.NumberFormat = "@"

# GENERAL/NUMBER columns (for formulas)
general_columns = ["B", "G"]
for col in general_columns:
    dst_ws.range(f"{col}:{col}").api.NumberFormat = "General"

# -----------------------------
# Paste columns B–I → D–K
# -----------------------------
dst_ws.range("D22").value = [row[:8] for row in data]

# -----------------------------
# Paste column K → column B
# -----------------------------
col_k_only = [[row[8]] for row in data]
dst_ws.range("B22").value = col_k_only

dst_wb.save()

print("Data copied successfully.")

# -----------------------------
# VALIDATION CHECK
# -----------------------------
print("\n===== VALIDATION RESULTS =====")

errors = []

src_users = [str(src_ws[f"H{r}"].value) for r in range(2, last_row + 1)]
src_targets = [str(src_ws[f"K{r}"].value) for r in range(2, last_row + 1)]

dst_users = [str(v) for v in dst_ws.range(f"J22:J{21+len(data)}").value]
dst_targets = [str(v) for v in dst_ws.range(f"B22:B{21+len(data)}").value]

def normalize_target(v):
    if v is None:
        return ""
    v = str(v).strip()
    try:
        return float(v)
    except:
        return v

for i, (su, st, du, dt) in enumerate(zip(src_users, src_targets, dst_users, dst_targets), start=22):

    match_user = (su.strip() == du.strip())

    st_n = normalize_target(st)
    dt_n = normalize_target(dt)
    match_target = (st_n == dt_n)

    status = "✔ MATCH" if (match_user and match_target) else "✖ MISMATCH"

    print(f"Row {i}: Source(H={su}, K={st}) → Dest(J={du}, B={dt})   {status}")

    if not (match_user and match_target):
        errors.append(i)

if errors:
    print("\n⚠️ MISMATCHES FOUND IN ROWS:", errors)
else:
    print("\n✔ All rows match correctly.")

print("Excel remains open.")
