import pandas as pd
import os
from openpyxl import load_workbook, Workbook

# Input file
input_file = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\dealers\Prima_Sales_Registration.xlsx"

# Output folder
output_folder = r"\\virgo\Work\OpControl\2016\Bogdan\Prima Sales Registration\dealers"

# WHSNAME groups
groups = [
    "KANDYLTD",
    "VARDAROV",
    "YONI-90",
    "MOBIMAKSLTD",
    "POSTBRANCH",
    "HAIKOVLTD",
    "ET_BOMA",
    "PREPAID TEAM"
]

sheet_name = "Prima_Sales_Registration_export"

# ---------- 1) FORMAT MAIN FILE ----------
print("Applying formats in main file...")

wb = load_workbook(input_file)
ws = wb[sheet_name]

date_format = "yyyy-mm-dd"
text_format = "@"

for row in range(2, ws.max_row + 1):  # skip header
    # A and H → date
    ws.cell(row=row, column=1).number_format = date_format
    ws.cell(row=row, column=8).number_format = date_format

    # B and C → force TEXT
    cell_b = ws.cell(row=row, column=2)
    cell_c = ws.cell(row=row, column=3)

    if cell_b.value is not None:
        cell_b.value = str(cell_b.value)
    if cell_c.value is not None:
        cell_c.value = str(cell_c.value)

    cell_b.number_format = text_format
    cell_c.number_format = text_format

wb.save(input_file)
print("✔ Main file formats applied and saved.\n")

# ---------- 2) LOAD DATA FOR SPLITTING ----------
df = pd.read_excel(input_file, sheet_name=sheet_name)

# Column F is the 6th column → index 5
column_name = df.columns[5]
print(f"Filtering based on column: {column_name}")

# ---------- 3) CREATE GROUP FILES WITH SAME FORMATTING ----------
for group in groups:
    print(f"Creating file for group: {group}")

    filtered_df = df[df[column_name].astype(str).str.contains(group, case=False, na=False)]

    if filtered_df.empty:
        print(f"  ⚠ No rows found for {group}, skipping")
        continue

    output_file = os.path.join(
        output_folder,
        f"Prima Sales Registration_{group}.xlsx"
    )

    # Create new workbook for this group
    new_wb = Workbook()
    ws_new = new_wb.active
    ws_new.title = "Data"

    # Write header
    ws_new.append(list(filtered_df.columns))

    # Write data rows
    for _, row in filtered_df.iterrows():
        ws_new.append(list(row))

    # Apply formats in new file
    max_row = ws_new.max_row

    for r in range(2, max_row + 1):
        # A and H → date
        ws_new.cell(row=r, column=1).number_format = date_format
        ws_new.cell(row=r, column=8).number_format = date_format

        # B and C → force TEXT
        cell_b = ws_new.cell(row=r, column=2)
        cell_c = ws_new.cell(row=r, column=3)

        if cell_b.value is not None:
            cell_b.value = str(cell_b.value)
        if cell_c.value is not None:
            cell_c.value = str(cell_c.value)

        cell_b.number_format = text_format
        cell_c.number_format = text_format

    new_wb.save(output_file)
    print(f"  ✔ Saved: {output_file}")

print("\n✅ Main file formatted and all group files created with the requested formats.")
